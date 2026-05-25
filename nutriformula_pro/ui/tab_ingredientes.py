from PyQt6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QTableWidget,
    QTableWidgetItem, QLineEdit, QComboBox, QPushButton,
    QHeaderView, QMessageBox, QFileDialog
)
from PyQt6.QtCore import Qt, QSortFilterProxyModel
from PyQt6.QtGui import QColor

from ui.dialogs import DialogoIngrediente


COLUMNAS = [
    'id', 'Nombre', 'Proteína%', 'EM Kcal', 'Fibra%', 'Grasa%',
    'Calcio%', 'Fósforo%', 'Lisina%', 'Metionina%', 'Colina mg/kg',
    'Precio/kg', 'Categoría'
]

CAMPOS_NUTRICIONALES = [
    'nombre', 'proteina', 'em_kcal', 'fibra', 'grasa',
    'calcio', 'fosforo', 'lisina', 'metionina', 'colina_mgr',
    'precio_kg', 'categoria'
]

CATEGORIAS = [
    'Todas', 'Cereales', 'Proteínas', 'Grasas',
    'Subproductos', 'Fibras', 'Aditivos', 'General'
]


class TabIngredientes(QWidget):
    def __init__(self, db, parent=None):
        super().__init__(parent)
        self.db = db
        self.insumos = []

        layout = QVBoxLayout(self)

        filtros_layout = QHBoxLayout()

        self.busqueda = QLineEdit()
        self.busqueda.setPlaceholderText('Buscar ingrediente...')
        self.busqueda.textChanged.connect(self._filtrar)
        filtros_layout.addWidget(self.busqueda)

        self.filtro_categoria = QComboBox()
        self.filtro_categoria.addItems(CATEGORIAS)
        self.filtro_categoria.currentTextChanged.connect(self._filtrar)
        filtros_layout.addWidget(self.filtro_categoria)

        layout.addLayout(filtros_layout)

        self.tabla = QTableWidget()
        self.tabla.setColumnCount(len(COLUMNAS) - 1)
        self.tabla.setHorizontalHeaderLabels(COLUMNAS[1:])
        self.tabla.setSelectionBehavior(
            QTableWidget.SelectionBehavior.SelectRows
        )
        self.tabla.setEditTriggers(
            QTableWidget.EditTrigger.NoEditTriggers
        )
        self.tabla.setAlternatingRowColors(True)
        self.tabla.horizontalHeader().setSectionResizeMode(
            QHeaderView.ResizeMode.Stretch
        )
        self.tabla.horizontalHeader().setSectionResizeMode(
            1, QHeaderView.ResizeMode.ResizeToContents
        )
        self.tabla.doubleClicked.connect(self._editar_seleccionado)
        layout.addWidget(self.tabla)

        botones_layout = QHBoxLayout()
        botones_layout.addStretch()

        btn_agregar = QPushButton('+ Agregar')
        btn_agregar.setObjectName('btn_primario')
        btn_agregar.clicked.connect(self._agregar)

        btn_editar = QPushButton('✏ Editar')
        btn_editar.clicked.connect(self._editar_seleccionado)

        btn_eliminar = QPushButton('🗑 Eliminar')
        btn_eliminar.setObjectName('btn_peligro')
        btn_eliminar.clicked.connect(self._eliminar)

        btn_importar = QPushButton('📥 Importar Excel')
        btn_importar.clicked.connect(self._importar_excel)

        btn_exportar = QPushButton('📤 Exportar Excel')
        btn_exportar.clicked.connect(self._exportar_excel)

        botones_layout.addWidget(btn_agregar)
        botones_layout.addWidget(btn_editar)
        botones_layout.addWidget(btn_eliminar)
        botones_layout.addWidget(btn_importar)
        botones_layout.addWidget(btn_exportar)

        layout.addLayout(botones_layout)

        self._cargar_datos()

    def _cargar_datos(self):
        self.insumos = self.db.get_insumos()
        self._filtrar()

    def _filtrar(self):
        texto = self.busqueda.text().lower()
        categoria = self.filtro_categoria.currentText()

        self.tabla.setRowCount(0)

        for insumo in self.insumos:
            if categoria != 'Todas' and insumo['categoria'] != categoria:
                continue

            if texto and texto not in insumo['nombre'].lower():
                continue

            fila = self.tabla.rowCount()
            self.tabla.insertRow(fila)

            valores = [
                insumo['nombre'],
                f"{insumo['proteina']:.2f}",
                f"{insumo['em_kcal']:.0f}",
                f"{insumo['fibra']:.2f}",
                f"{insumo['grasa']:.2f}",
                f"{insumo['calcio']:.4f}",
                f"{insumo['fosforo']:.4f}",
                f"{insumo['lisina']:.2f}",
                f"{insumo['metionina']:.2f}",
                f"{insumo['colina_mgr']:.0f}",
                f"${insumo['precio_kg']:.4f}",
                insumo['categoria'],
            ]

            for col, valor in enumerate(valores):
                item = QTableWidgetItem(valor)
                item.setData(Qt.ItemDataRole.UserRole, insumo['id'])
                self.tabla.setItem(fila, col, item)

    def _get_insumo_id_seleccionado(self):
        seleccion = self.tabla.selectionModel().selectedRows()
        if not seleccion:
            return None
        fila = seleccion[0].row()
        item = self.tabla.item(fila, 0)
        if item:
            return item.data(Qt.ItemDataRole.UserRole)
        return None

    def _agregar(self):
        dialogo = DialogoIngrediente(self.db, parent=self)
        if dialogo.exec() == DialogoIngrediente.DialogCode.Accepted:
            self._cargar_datos()

    def _editar_seleccionado(self):
        insumo_id = self._get_insumo_id_seleccionado()
        if not insumo_id:
            return
        dialogo = DialogoIngrediente(self.db, insumo_id, parent=self)
        if dialogo.exec() == DialogoIngrediente.DialogCode.Accepted:
            self._cargar_datos()

    def _eliminar(self):
        insumo_id = self._get_insumo_id_seleccionado()
        if not insumo_id:
            QMessageBox.warning(self, 'Atención', 'Selecciona un ingrediente')
            return

        resp = QMessageBox.question(
            self, 'Confirmar',
            '¿Eliminar este ingrediente?',
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No
        )
        if resp == QMessageBox.StandardButton.Yes:
            try:
                self.db.eliminar_insumo(insumo_id)
                self._cargar_datos()
            except Exception as e:
                QMessageBox.critical(self, 'Error', str(e))

    def _importar_excel(self):
        try:
            import openpyxl
        except ImportError:
            QMessageBox.warning(
                self, 'Error', 'openpyxl no está instalado'
            )
            return

        archivo, _ = QFileDialog.getOpenFileName(
            self, 'Importar Excel', '', 'Excel Files (*.xlsx *.xls)'
        )
        if not archivo:
            return

        try:
            wb = openpyxl.load_workbook(archivo)
            ws = wb.active

            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row or not row[0]:
                    continue
                datos = {
                    'nombre': str(row[0]),
                    'proteina': float(row[1]) if row[1] else 0,
                    'em_kcal': float(row[2]) if row[2] else 0,
                    'fibra': float(row[3]) if row[3] else 0,
                    'grasa': float(row[4]) if row[4] else 0,
                    'calcio': float(row[5]) if row[5] else 0,
                    'fosforo': float(row[6]) if row[6] else 0,
                    'lisina': float(row[7]) if row[7] else 0,
                    'metionina': float(row[8]) if row[8] else 0,
                    'colina_mgr': float(row[9]) if row[9] else 0,
                    'precio_kg': float(row[10]) if row[10] else 0,
                    'categoria': str(row[11]) if row[11] else 'General',
                }
                try:
                    self.db.insertar_insumo(datos)
                except Exception:
                    pass

            self._cargar_datos()
            QMessageBox.information(
                self, 'Éxito', 'Ingredientes importados correctamente'
            )
        except Exception as e:
            QMessageBox.critical(self, 'Error', f'Error al importar: {str(e)}')

    def _exportar_excel(self):
        try:
            import openpyxl
        except ImportError:
            QMessageBox.warning(
                self, 'Error', 'openpyxl no está instalado'
            )
            return

        archivo, _ = QFileDialog.getSaveFileName(
            self, 'Exportar Excel', '', 'Excel Files (*.xlsx)'
        )
        if not archivo:
            return

        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = 'Ingredientes'

            encabezados = COLUMNAS[1:]
            ws.append(encabezados)

            for insumo in self.insumos:
                ws.append([
                    insumo['nombre'],
                    insumo['proteina'],
                    insumo['em_kcal'],
                    insumo['fibra'],
                    insumo['grasa'],
                    insumo['calcio'],
                    insumo['fosforo'],
                    insumo['lisina'],
                    insumo['metionina'],
                    insumo['colina_mgr'],
                    insumo['precio_kg'],
                    insumo['categoria'],
                ])

            wb.save(archivo)
            QMessageBox.information(
                self, 'Éxito', f'Exportado a:\n{archivo}'
            )
        except Exception as e:
            QMessageBox.critical(self, 'Error', f'Error al exportar: {str(e)}')
